"""
Turns a detailed_signals_*.xlsx into a stored collection.

Each collection is stored twice:
  snapshots.data_gz   the whole workbook as gzip JSON {sheet: rows} — the dashboard reads this as-is,
                      so new sheets from the scanner show up with no database change
  snapshot_stocks     one row per stock (price, change, bias, signals, health) for fast comparisons
"""
import gzip
import json
import math
import os
import re
from collections import Counter, defaultdict
from datetime import datetime

from openpyxl import load_workbook

from .db import pool


def _clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return ""
    if isinstance(v, datetime):
        return v.isoformat()
    return v


def workbook_to_sheets(path):
    """Same shape as XLSX.utils.sheet_to_json(ws, {defval: ""}) in the browser."""
    wb = load_workbook(path, read_only=True, data_only=True)
    out = {}
    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            try:
                header = next(rows)
            except StopIteration:
                out[ws.title] = []
                continue
            keys, seen = [], Counter()
            for i, h in enumerate(header):
                k = str(h).strip() if h is not None and str(h).strip() else f"__EMPTY_{i}"
                if seen[k]:
                    k = f"{k}_{seen[k]}"
                seen[k] += 1
                keys.append(k)
            data = []
            for r in rows:
                if r is None or all(v is None for v in r):
                    continue
                data.append({keys[i]: _clean(r[i]) if i < len(r) else "" for i in range(len(keys))})
            out[ws.title] = data
    finally:
        wb.close()
    return out


def _num(v):
    try:
        f = float(v)
        return None if math.isnan(f) else f
    except (TypeError, ValueError):
        return None


def stock_rows(sheets):
    """One summary row per stock for snapshot_stocks."""
    stocks = {}

    def get(sym):
        sym = str(sym).replace("NSE:", "").replace(",", "").strip().upper()   # same as cleanSym in the dashboard
        if not sym:
            return None
        return stocks.setdefault(sym, {"symbol": sym, "price": None, "change_pct": None, "sector": None,
                                       "industry": None, "bias": None, "signals": [], "health": None, "is_fno": False})

    for r in sheets.get("Master_Stock_Data", []):
        s = get(r.get("Symbol", ""))
        if s:
            s["price"] = _num(r.get("Price"))
            s["change_pct"] = _num(r.get("Change_Pct"))
            s["sector"] = r.get("Sector") or None
            s["industry"] = r.get("Industry") or None
    for r in sheets.get("Zone_Levels", []):
        s = get(r.get("Symbol", ""))
        if s:
            if s["price"] is None:
                s["price"] = _num(r.get("Price"))
            if s["change_pct"] is None:
                s["change_pct"] = _num(r.get("Change_Pct"))
            s["sector"] = s["sector"] or r.get("Sector") or None
            s["industry"] = s["industry"] or r.get("Industry") or None
            s["is_fno"] = s["is_fno"] or r.get("Is_FNO") == "Yes"
    bias_votes = defaultdict(Counter)
    for r in sheets.get("Flat_Data_For_Slicers", []):
        s = get(r.get("Symbol", ""))
        if s:
            s["signals"].append(str(r.get("Signal_Name", "")))
            bias_votes[s["symbol"]][str(r.get("Trading_Bias", ""))] += 1
            s["is_fno"] = s["is_fno"] or r.get("Is_FNO") == "Yes"
    for r in sheets.get("Top_Opportunities", []):
        s = get(r.get("Symbol", ""))
        if s and r.get("Trading_Bias"):
            s["bias"] = r.get("Trading_Bias")
    for sym, votes in bias_votes.items():
        if not stocks[sym]["bias"]:
            stocks[sym]["bias"] = votes.most_common(1)[0][0] or None
    for r in sheets.get("Price_Health", []):
        s = get(r.get("Symbol", ""))
        if s:
            s["health"] = r.get("Health") or None
    return list(stocks.values())


def save_snapshot(sheets, taken_at, source, slot=None, label=None, excel_path=None):
    raw = json.dumps(sheets, separators=(",", ":"), ensure_ascii=False).encode("utf-8")
    gz = gzip.compress(raw, compresslevel=6)
    rows = stock_rows(sheets)
    n_signals = len(sheets.get("Flat_Data_For_Slicers", []))
    n_stocks = len({r["symbol"] for r in rows if r["signals"]}) or len(rows)
    label = label or taken_at.strftime("%d %b %Y %H:%M")
    with pool.connection() as c:
        with c.transaction():
            snap = c.execute(
                "INSERT INTO snapshots (taken_at, trade_date, slot, source, label, stocks, signals, sheets, data_gz, size_bytes, excel_path) "
                "VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s) RETURNING id",
                (taken_at, taken_at.date(), slot, source, label, n_stocks, n_signals, list(sheets.keys()),
                 gz, len(gz), excel_path)).fetchone()
            sid = snap["id"]
            with c.cursor().copy("COPY snapshot_stocks (snapshot_id, symbol, price, change_pct, sector, industry, bias, "
                                 "n_signals, health, is_fno, signals) FROM STDIN") as cp:
                for r in rows:
                    cp.write_row((sid, r["symbol"], r["price"], r["change_pct"], r["sector"], r["industry"], r["bias"],
                                  len(r["signals"]), r["health"], r["is_fno"], "|".join(r["signals"])))
    return sid


FILE_DATE_RE = re.compile(r"(20\d{2})(\d{2})(\d{2})(?:[_-]?(\d{2})(\d{2})(\d{2})?)?")


def taken_at_from_filename(name, tz, default_time=(15, 30)):
    """detailed_signals_20260916.xlsx -> 16 Sep 2026 15:30 ; ..._20260916_0945.xlsx -> 09:45"""
    m = FILE_DATE_RE.search(os.path.basename(name))
    if not m:
        return None
    y, mo, d, hh, mm = int(m[1]), int(m[2]), int(m[3]), m[4], m[5]
    h, mi = (int(hh), int(mm)) if hh and mm and int(hh) < 24 and int(mm) < 60 else default_time
    try:
        return datetime(y, mo, d, h, mi, tzinfo=tz)
    except ValueError:
        return None
